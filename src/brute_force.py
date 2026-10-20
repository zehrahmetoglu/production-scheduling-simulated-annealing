import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from collections import defaultdict
import matplotlib.pyplot as plt
import random
import itertools
import time
random.seed(100)
class FactoryManager:
    def __init__(self):
        self.jobs = {}
        self.machines = {}
    def generate_jobs(self, num_jobs):
        for job_id in range(1, num_jobs + 1):
            self.jobs[job_id] = {"processing_times": []}
    def load_processing_times(self, filename):
        df = pd.read_excel(filename, index_col='ID')
        for job_id, row in df.iterrows():
            processing_times = row.dropna().drop('DueDate').tolist()
            due_date = row['DueDate'] if pd.notna(row['DueDate']) else None
            self.jobs[job_id] = { "stages": processing_times, "due_date": due_date }
    def load_machine_speeds(self, filename, scenario):
        df = pd.read_excel(filename, sheet_name=scenario)
        for index, row in df.iterrows():
            machine_name = row['Machine'].strip()
            stage_speeds = {}
            for count in range(1, 7):
                stage_key = f'stage{count}'
                if f'Stage{count}' in row:
                    stage_speeds[stage_key] = row[f'Stage{count}']
                else:
                    break
            self.machines[machine_name] = stage_speeds
    def create_joblist(self):
        joblist = {}
        for job in self.jobs:
            if job not in joblist:
                joblist[job] = {}
            for machine_count, machine in enumerate(self.machines):
                for stage_count, processing_time in enumerate(self.jobs[job]["stages"]):
                    true_processing_time = processing_time * self.machines[f"Machine{machine_count + 1}"][ f"stage{stage_count + 1}"]
                    if stage_count not in joblist[job]:
                        joblist[job][stage_count] = {}
                    joblist[job][stage_count][machine] = true_processing_time
        return joblist, self.machines
class SPTLPTEDD:
    def __init__(self, excel_path, choice):
        self.excel_path = excel_path
        self.jobs_data = {}
        self.stage1_times = []
        self.due_dates = []
        self.user_initial_solution_choice = choice
    def load_excel_data(self):
        df = pd.read_excel(self.excel_path, sheet_name=0)
        self.stage1_times = df['Stage1'].tolist()
        self.due_dates = df['DueDate'].tolist()
    def set_jobs_data(self, jobs_data):
        self.jobs_data = jobs_data
    def reorder_jobs(self):
        if not self.jobs_data or not self.stage1_times:
            raise ValueError("Jobs data or Stage1 times are not loaded properly.")
        if self.user_initial_solution_choice == 1:
            sorted_jobs = sorted(self.jobs_data.items(), key=lambda x: self.stage1_times[x[0] - 1])
        else:
            sorted_jobs = sorted(self.jobs_data.items(), key=lambda x: self.stage1_times[x[0] - 1], reverse=True)
        sorted_data = {job[0]: job[1] for job in sorted_jobs}
        return sorted_data
    def reorder_jobs_by_due_date(self):
        if not self.jobs_data or not self.due_dates:
            raise ValueError("Jobs data or DueDate times are not loaded properly.")
        # Sort jobs by DueDate in ascending order
        sorted_jobs = sorted(self.jobs_data.items(), key=lambda x: self.due_dates[x[0] - 1])
        return {job[0]: job[1] for job in sorted_jobs}
class Random:
    def __init__(self, job_data):
        self.job_data = job_data
    def reorder_jobs(self):
        # Extract job keys and shuffle them randomly
        job_keys = list(self.job_data.keys())
        random.shuffle(job_keys)
        # Create a new dictionary to store reordered jobs
        new_job_order = {}
        # Assign the old job data to the new job keys in the shuffled order
        for old_key in job_keys:
            new_job_order[old_key] = self.job_data[old_key]
        return new_job_order
class Swap:
    def __init__(self, data):
        self.data = data
    def swap_jobs(self):
        # Get all job IDs
        job_ids = list(self.data.keys())
        # Randomly select two distinct jobs to swap
        job1, job2 = random.sample(job_ids, 2)
        # Swap the jobs including their keys to reflect the change in positions
        temp = self.data[job1]
        self.data[job1] = self.data[job2]
        self.data[job2] = temp
        # Now update the keys
        new_data = {}
        for key in self.data:
            if key == job1:
                new_data[job2] = self.data[key]
            elif key == job2:
                new_data[job1] = self.data[key]
            else:
                new_data[key] = self.data[key]
        self.data = new_data
        return self.data
class Insert:
    def __init__(self, data):
        self.data = data
        self.keys_order = list(data.keys())  # Dynamically get keys from the dictionary
        self.new_dict = {}
    def random_reorder(self):
        # If there's only one key, no reorder is needed
        if len(self.keys_order) <= 1:
            self.new_dict = self.data
            return
        # Copy the current key order to manipulate
        temp_keys_order = self.keys_order[:]
        # Select a random key and insert it at a new random position
        key_to_move = random.choice(temp_keys_order)
        temp_keys_order.remove(key_to_move)
        new_position = random.randint(0, len(temp_keys_order))
        temp_keys_order.insert(new_position, key_to_move)
        # Update the original order to reflect this change
        self.keys_order = temp_keys_order
        # Create a new dictionary based on the new order of keys
        self.new_dict = {}
        for key in self.keys_order:
            self.new_dict[key] = self.data[key]
    def get_reordered_dict(self):
        return self.new_dict
class AdjacentSwap:
    def __init__(self, data):
        self.data = data
        self.keys = list(data.keys())  # Store the keys in the order they are initially provided
    def random_swap_and_reconstruct(self):
        # Perform a random adjacent swap on the list of keys
        if len(self.keys) > 1:  # Ensure there's more than one key to swap
            swap_index = random.randint(0, len(self.keys) - 2)  # Select a random index for swapping
            # Swap the current index with the next one
            self.keys[swap_index], self.keys[swap_index + 1] = self.keys[swap_index + 1], self.keys[swap_index]
        # Reconstruct the dictionary with the new keys order
        new_data = {key: self.data[key] for key in self.keys}
        self.data = new_data  # Update the internal data with the newly ordered data
    def get_data(self):
        return self.data
class Dispatch:
    def __init__(self, jobs_data, use_random=False):
        self.jobs_data = jobs_data
        self.machine_workload = {}
        self.scheduled_jobs = {}
        self.use_random = use_random
        self.initialize_workloads()
    def initialize_workloads(self):
        for job in self.jobs_data:
            for stage in self.jobs_data[job]:
                if stage not in self.machine_workload:
                    self.machine_workload[stage] = {}
                for machine in self.jobs_data[job][stage]:
                    if machine not in self.machine_workload[stage]:
                        self.machine_workload[stage][machine] = 0
    def find_min_workload_machine(self, stage, job_stage_data):
        if self.use_random:
            machine = random.choice(list(job_stage_data.keys()))
            new_workload = self.machine_workload[stage][machine] + job_stage_data[machine]
            return machine, new_workload
        else:
            min_workload = float('inf')
            selected_machine = None
            for machine, processing_time in job_stage_data.items():
                new_workload = self.machine_workload[stage].get(machine, 0) + processing_time
                if new_workload < min_workload:
                    min_workload = new_workload
                    selected_machine = machine
            return selected_machine, min_workload
    def schedule_jobs(self):
        for job_id, stages in self.jobs_data.items():
            self.scheduled_jobs[job_id] = {}
            for stage_id, machines in stages.items():
                machine, new_workload = self.find_min_workload_machine(stage_id, machines)
                self.scheduled_jobs[job_id][stage_id] = {machine: machines[machine]}
                self.machine_workload[stage_id][machine] += machines[machine]
    def get_scheduled_jobs(self):
        return self.scheduled_jobs
    def get_machine_workloads(self):
        return self.machine_workload
    def swap_random_jobs(self, new_jobs_data):
        self.jobs_data = new_jobs_data
        self.initialize_workloads()
        self.schedule_jobs()
        job_ids = list(self.jobs_data.keys())
        if len(job_ids) > 1:
            job1, job2 = random.sample(job_ids, 2)
            self.scheduled_jobs[job1], self.scheduled_jobs[job2] = self.scheduled_jobs[job2], self.scheduled_jobs[job1]
            self.initialize_workloads()
            self.schedule_jobs()
        return self.get_scheduled_jobs()
class Formatter1:
    def __init__(self):
        self.jobs = {}
    def add_job_data(self, job_data):
        # Adds or updates job data
        self.jobs.update(job_data)
    def get_formatted_stage_dictionary(self):
        stage_processing_times = defaultdict(lambda: defaultdict(dict))
        for job, stages in self.jobs.items():
            for stage, machine_info in stages.items():
                for machine, time in machine_info.items():
                    if job not in stage_processing_times[stage][machine]:
                        stage_processing_times[stage][machine][job] = time
        return self._convert_to_regular_dict(stage_processing_times)
    def _convert_to_regular_dict(self, d):
        if isinstance(d, defaultdict):
            d = {k: self._convert_to_regular_dict(v) for k, v in d.items()}
        return d
class Formatter2:
    def __init__(self, existing_data):
        self.schedule = existing_data
    def load_due_dates_from_excel(self, file):
        df = pd.read_excel(file, sheet_name=0)
        due_dates = df.iloc[:, -1].to_dict()
        job_ids = df.index.to_list()
        job_due_dates = {job_id + 1: due_dates[i] for i, job_id in enumerate(job_ids)}
        for stage, machines in self.schedule.items():
            for machine, jobs in machines.items():
                for job_id, details in list(jobs.items()):
                    if job_id in job_due_dates:
                        if isinstance(details, tuple):
                            current_due_date = details[1]
                            if current_due_date != job_due_dates[job_id]:
                                print(f"Updating due date for job {job_id} from {current_due_date}" f" to {job_due_dates[job_id]}")
                            jobs[job_id] = [details[0], job_due_dates[job_id]]
                        else:
                            jobs[job_id] = [details, job_due_dates[job_id]]
    def get_schedule(self):
        return self.schedule
class ScheduleProcessor1:
    def __init__(self, data_to_sort):
        self.data_to_sort = data_to_sort
        self.schedule = {}
        self.job_end_times = {}
        self.machine_end_times = {}
    def schedule_jobs(self):
        for stage_count, stage in enumerate(sorted(self.data_to_sort.keys())):
            self.machine_end_times[stage] = {}
            self.schedule[stage] = {}
            for machine in self.data_to_sort[stage]:
                self.schedule[stage][machine] = {}
                if machine not in self.machine_end_times[stage]:
                    self.machine_end_times[stage][machine] = 0
                for job, (processing_time, _) in self.data_to_sort[stage][machine].items():
                    if job not in self.job_end_times:
                        self.job_end_times[job] = 0
                    start_time = max(self.job_end_times[job], self.machine_end_times[stage][machine])
                    end_time = start_time + processing_time
                    self.job_end_times[job] = end_time
                    self.machine_end_times[stage][machine] = end_time
                    self.schedule[stage][machine][job] = [start_time, end_time]
    def get_schedule(self):
        return self.schedule
class ScheduleProcessor2:
    def __init__(self, schedule_process_data, excel_path_for_tardiness):
        self.schedule_data = schedule_process_data
        self.excel_path_tardiness = excel_path_for_tardiness
    def makespan(self):
        last_stage_key = max(self.schedule_data.keys())
        makespan = 0
        for machine in self.schedule_data[last_stage_key]:
            job_timings = self.schedule_data[last_stage_key][machine].values()
            for timing in job_timings:
                makespan = max(makespan, timing[1])
        return makespan
    def total_completion_time(self):
        last_stage = max(self.schedule_data.keys())
        total_completion_time = 0
        for machine, jobs in self.schedule_data[last_stage].items():
            for job, times in jobs.items():
                total_completion_time += times[1]
        return total_completion_time
    def calculate_tardiness(self):
        due_dates = pd.read_excel(self.excel_path_tardiness, sheet_name=0)
        due_dates = due_dates['DueDate'].tolist()
        completion_times = {}
        for stage in self.schedule_data.values():
            for machine, jobs in stage.items():
                for job, times in jobs.items():
                    completion_time = times[1]
                    if job in completion_times:
                        if completion_time > completion_times[job]:
                            completion_times[job] = completion_time
                    else:
                        completion_times[job] = completion_time
        tardiness = {}
        for job, completion_time in completion_times.items():
            if completion_time - due_dates[job - 1] > 0:
                tardiness[job] = completion_time - due_dates[job - 1]
        total_tardiness = sum(tardiness.values())
        return tardiness, total_tardiness
    def process_schedule(self):
        makespan = self.makespan()
        total_completion_time = self.total_completion_time()
        total_tardiness = self.calculate_tardiness()[1]
        parameters = { "makespan": makespan, "total tardiness": total_tardiness, "total completion time": total_completion_time }
        return parameters
class ExcelExporter:
    def __init__(self, data, filename, sheet_name='Result Schedule'):
        self.data = data
        self.filename = filename
        self.sheet_name = sheet_name
    def export_to_excel(self):
        rows = []
        for stage, machines in self.data.items():
            for machine, jobs in machines.items():
                for job, times in jobs.items():
                    rows.append({'Stage': stage, 'Machine': machine, 'Job': job, 'Enter Time': times[0], 'Leave Time': times[1]})
        df = pd.DataFrame(rows)
        df.sort_values(by='Job', ascending=True, inplace=True)
        book = load_workbook(self.filename)
        if self.sheet_name in book.sheetnames:
            del book[self.sheet_name]
        ws = book.create_sheet(self.sheet_name)
        ws.sheet_state = 'visible'
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
            for c_idx, number in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=number)
        book.save(self.filename)
        print(f"Data exported to {self.filename} into sheet '{self.sheet_name}'.")
class Grapher:
    def __init__(self, data):
        self.data = data
        self.sorted_temperatures = []
        self.sorted_costs = []
        self._sort_data()
    def _sort_data(self):
        temperatures = self.data['iteration']
        costs = self.data['cost']
        sorted_indices = sorted(range(len(temperatures)), key=lambda i: temperatures[i], reverse=True)
        self.sorted_temperatures = [temperatures[i] for i in sorted_indices]
        self.sorted_costs = [costs[i] for i in sorted_indices]
    def plot(self):
        plt.figure(figsize=(10, 5))
        plt.plot(self.sorted_temperatures, self.sorted_costs, marker='o', markersize=0.1)
        plt.title('Cost vs. Temperature')
        plt.xlabel('Iteration')
        plt.ylabel('Cost')
        plt.ylim(min(self.sorted_costs) - 10, max(self.sorted_costs) + 10)
        plt.grid(True)
        min_cost_index = self.sorted_costs.index(min(self.sorted_costs))
        min_temperature = self.sorted_temperatures[min_cost_index]
        min_cost = self.sorted_costs[min_cost_index]
        plt.scatter([min_temperature], [min_cost], color='red', s=50)
        plt.show()
class DataExporter:
    def __init__(self, data, file_path_to_export='outputs/experiment_results.xlsx'):
        self.data = data
        self.file_path = file_path_to_export
    def _parse_data(self):
        rows = []
        for scenario, scenario_data in self.data.items():
            for method, method_data in scenario_data.items():
                for operation, operation_data in method_data.items():
                    for metric, metric_data in operation_data.items():
                        for solution_type, values in metric_data.items():
                            row = { 'Scenario': scenario, 'Method': method, 'Operation': operation, 'Metric': metric, 'Solution Type': solution_type, 'Makespan': values.get('makespan', 'N/A'), 'Total Tardiness': values.get('total tardiness', 'N/A'), 'Total Completion Time': values.get('total completion time', 'N/A') }
                            rows.append(row)
        return rows
    def export_to_excel(self):
        rows = self._parse_data()
        df = pd.DataFrame(rows)
        column_order = ['Scenario', 'Method', 'Operation', 'Metric', 'Solution Type', 'Makespan', 'Total Tardiness', 'Total Completion Time']
        df = df[column_order]
        with pd.ExcelWriter(self.file_path, engine='xlsxwriter') as writer:
            df.to_excel(writer, index=False, sheet_name='Results')
        return f"Excel file created successfully at {self.file_path}"
class PermutationGenerator:
    def __init__(self, data):
        self.data = data
    def create_permutation_dicts(self):
        # Extract the job order from the keys of the data dictionary
        job_order = list(self.data.keys())
        # Generate all possible permutations of job orders
        permutations = list(itertools.permutations(job_order))
        # Initialize a dictionary to store all possible job sequences
        permutation_dicts = {}
        # Assign each permutation to a new dictionary entry
        for i, perm in enumerate(permutations):
            sequence_data = {}
            # For each job in the permutation, match the job key to its position in the sequence
            for job_index, job_key in enumerate(perm):
                sequence_data[job_key] = self.data[job_key]
            permutation_dicts[i] = sequence_data
        return permutation_dicts
class Timer:
    def __init__(self):
        self.start_time = None
        self.end_time = None
    def start(self):
        self.start_time = time.time()
    def stop(self):
        self.end_time = time.time()
    def elapsed_time_seconds(self):
        if self.start_time is None:
            return 0
        if self.end_time is None:
            return time.time() - self.start_time
        return round(self.end_time - self.start_time, 2)
    def reset(self):
        self.start_time = None
        self.end_time = None
if __name__ == "__main__":
    file_path = 'data/sample_jobs_small.xlsx'
    manager = FactoryManager()
    manager.generate_jobs(1)
    manager.load_processing_times(file_path)
    setup_scenario = int(input("Set up scenario: "))
    manager.load_machine_speeds(file_path, setup_scenario)
    raw_data, machine_data = manager.create_joblist()
    job_sequence_generator = PermutationGenerator(raw_data)
    permutation_data = job_sequence_generator.create_permutation_dicts()
    total_keys = len(permutation_data)
    user_criteria_choice = int(input("select success criteria\n" "( makespan / total tardiness / total completion time)" "\n(1 / 2 / 3) "))
    if user_criteria_choice == 1:
        user_criteria_choice = 'makespan'
    elif user_criteria_choice == 2:
        user_criteria_choice = 'total tardiness'
    elif user_criteria_choice == 3:
        user_criteria_choice = 'total completion time'
    else:
        user_criteria_choice = "none"
        print("illegal")
    select_best_order = {}
    positive_infinity = float('inf')
    best_order_values = {}
    best_order = {}
    timer = Timer()
    timer.start()
    for key_index, sequence_key in enumerate(permutation_data):
        sequence = permutation_data[sequence_key]
        scheduler = Dispatch(sequence)
        scheduler.schedule_jobs()
        assigned_data = scheduler.get_scheduled_jobs()
        formatter_1 = Formatter1()
        formatter_1.add_job_data(assigned_data)
        formatted_data = formatter_1.get_formatted_stage_dictionary()
        formatter_2 = Formatter2(formatted_data)
        formatter_2.load_due_dates_from_excel(file_path)
        formatted_data_2 = formatter_2.get_schedule()
        processor1 = ScheduleProcessor1(formatted_data_2)
        processor1.schedule_jobs()
        processed_order_1 = processor1.get_schedule()
        processor2 = ScheduleProcessor2(processed_order_1, file_path)
        process2 = processor2.process_schedule()
        processed_order_2 = process2
        if processed_order_2[user_criteria_choice] <= positive_infinity:
            positive_infinity = processed_order_2[user_criteria_choice]
            print(positive_infinity)
            best_order_values = processed_order_2
            best_order = processed_order_1
        else:
            best_order_values = best_order_values
            best_order = best_order
    print(best_order_values)
    print(best_order)
    timer.stop()
    print("Elapsed time:", timer.elapsed_time_seconds(), "seconds")
